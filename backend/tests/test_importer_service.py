import io
from types import SimpleNamespace

import openpyxl

from katalon.services.importer_service import (
    _cluster_values,
    apply_mapping,
    detect_delimiter,
    dry_run,
    parse_csv,
    parse_excel,
)


def _make_xlsx(headers: list[str], rows: list[list]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_parse_excel_basic() -> None:
    data = _make_xlsx(["title", "creator"], [["Foto 1", "Maier"], ["Foto 2", "Huber"]])
    headers, rows = parse_excel(data)
    assert headers == ["title", "creator"]
    assert len(rows) == 2
    assert rows[0]["title"] == "Foto 1"
    assert rows[1]["creator"] == "Huber"


def test_parse_excel_empty_sheet() -> None:
    wb = openpyxl.Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    headers, rows = parse_excel(buf.getvalue())
    assert headers == []
    assert rows == []


def test_detect_delimiter_comma() -> None:
    assert detect_delimiter("a,b,c\n1,2,3") == ","


def test_detect_delimiter_semicolon() -> None:
    assert detect_delimiter("a;b;c\n1;2;3") == ";"


def test_parse_csv_basic() -> None:
    content = b"title,creator\nFoto 1,Maier\nFoto 2,Huber\n"
    headers, rows = parse_csv(content)
    assert headers == ["title", "creator"]
    assert len(rows) == 2
    assert rows[0]["title"] == "Foto 1"


def test_apply_mapping() -> None:
    rows = [{"title": "Foto 1", "creator": "Maier"}]
    mapping = {"title": "title", "creator": "photographer"}
    records, idnos = apply_mapping(rows, mapping)
    # Without field_defs, fields are treated as non-repeatable -> plain string
    assert records[0]["title"] == "Foto 1"
    assert records[0]["photographer"] == "Maier"
    assert idnos[0] is None


def test_apply_mapping_with_split_transform_repeatable() -> None:
    """Split transform with is_repeatable=True should produce a list."""
    rows = [{"languages": "Deutsch, Englisch, Französisch"}]
    mapping = {"languages": {"target": "languages", "transforms": [{"type": "split", "delimiter": ",", "filter_empty": True}]}}
    # Mock field_defs with is_repeatable=True
    class MockField:
        field_type = "text"
        is_repeatable = True
    records, idnos = apply_mapping(rows, mapping, {"languages": MockField()})
    assert records[0]["languages"] == ["Deutsch", "Englisch", "Französisch"]


def test_apply_mapping_with_xml_repeated_element_list_value() -> None:
    """XmlFormat.parse() returns a list for repeated elements even when mapped
    to a non-repeatable field; apply_mapping must not crash on list values.
    Regression for AttributeError: 'list' object has no attribute 'strip'.
    """
    rows = [{"title": ["Foto 1", "Foto 1 (Alt)"], "idno_col": ["OBJ-1"]}]
    mapping = {"title": "title", "idno_col": "__idno__"}
    records, idnos = apply_mapping(rows, mapping)
    assert records[0]["title"] == "Foto 1"
    assert idnos[0] == "OBJ-1"


def test_apply_mapping_with_dict_target() -> None:
    rows = [{"title": "Foto 1"}]
    mapping = {"title": {"target": "title"}}
    records, idnos = apply_mapping(rows, mapping)
    assert records[0]["title"] == "Foto 1"


def test_apply_transforms_replace() -> None:
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("Hello World", [{"type": "replace", "search": "World", "replace": "Katalon", "case_sensitive": True}])
    assert result == ["Hello Katalon"]


def test_apply_transforms_regex_extract() -> None:
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("Jahr: 1996", [{"type": "regex_extract", "pattern": r"\d{4}", "group": 0}])
    assert result == ["1996"]


def test_apply_transforms_vocab_map() -> None:
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("DE", [{"type": "vocab_map", "vocab_map": {"DE": "Deutsch", "EN": "Englisch"}, "strict": False}])
    assert result == ["Deutsch"]


def test_apply_transforms_expression() -> None:
    from katalon.services.importer_service import apply_transforms
    # Jinja2 syntax
    result = apply_transforms("hello", [{"type": "expression", "expression": "{{ value | upper }}"}])
    assert result == ["HELLO"]
    # Backward compat: bare ${value} still works
    result2 = apply_transforms("world", [{"type": "expression", "expression": "prefix_${value}_suffix"}])
    assert result2 == ["prefix_world_suffix"]


def test_apply_transforms_pipeline() -> None:
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("  a, b, c  ", [
        {"type": "trim"},
        {"type": "split", "delimiter": ",", "filter_empty": True},
    ])
    assert result == ["a", "b", "c"]


def test_apply_transforms_vocab_map_strict_drops_unknown() -> None:
    """vocab_map with strict=True drops values not in the map."""
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("XX", [{"type": "vocab_map", "vocab_map": {"DE": "Deutsch"}, "strict": True}])
    assert result == []


def test_apply_transforms_vocab_map_strict_passes_known() -> None:
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("DE", [{"type": "vocab_map", "vocab_map": {"DE": "Deutsch"}, "strict": True}])
    assert result == ["Deutsch"]


def test_apply_transforms_regex_extract_invalid_pattern() -> None:
    """Invalid regex pattern falls back to original value instead of crashing."""
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("hello", [{"type": "regex_extract", "pattern": "[invalid("}])
    assert result == ["hello"]


def test_apply_transforms_expression_multiple_placeholders() -> None:
    """Jinja2 expression with value referenced multiple times."""
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("hi", [{"type": "expression", "expression": "{{ value }}-{{ value | upper }}"}])
    assert result == ["hi-HI"]


def test_apply_transforms_split_empty_delimiter_returns_original() -> None:
    """Empty delimiter is a no-op — returns the original value unchanged."""
    from katalon.services.importer_service import apply_transforms
    result = apply_transforms("a;b;c", [{"type": "split", "delimiter": ""}])
    assert result == ["a;b;c"]


def test_dry_run_no_errors() -> None:
    rows = [{"title": "Foto 1"}]
    mapping = {"title": "title"}
    result = dry_run(rows, mapping)
    assert result["total"] == 1
    assert result["valid"] == 1
    assert result["errors"] == []


def test_dry_run_empty_mapping() -> None:
    rows = [{"title": "Foto 1"}]
    mapping: dict = {}
    result = dry_run(rows, mapping)
    assert result["errors"][0]["row"] == 2  # row 1 is the header; first data row is row 2


def test_cluster_values_groups_typo_variants() -> None:
    counts = {"Berlin": 5, "berlin": 2, "Brlin": 1, "Hamburg": 3}
    clusters = _cluster_values(counts)
    assert len(clusters) == 1
    cluster = clusters[0]
    assert cluster["canonical"] == "Berlin"  # most frequent wins
    assert set(cluster["variants"]) == {"Berlin", "berlin", "Brlin"}
    assert cluster["counts"] == {"Berlin": 5, "berlin": 2, "Brlin": 1}


def test_cluster_values_no_singletons() -> None:
    counts = {"Berlin": 5, "Hamburg": 3, "Munich": 1}
    assert _cluster_values(counts) == []


def test_cluster_values_respects_max_cluster_size() -> None:
    counts = {f"aaaaaaaaa{i}": 1 for i in range(20)}
    clusters = _cluster_values(counts, threshold=0.5, max_cluster_size=4)
    assert all(len(c["variants"]) <= 4 for c in clusters)


def test_dry_run_vocab_clusters_suggests_canonical_value() -> None:
    rows = [
        {"place": "Berlin"},
        {"place": "Berlin"},
        {"place": "berlin"},
        {"place": "Brlin"},
    ]
    mapping = {"place": {"target": "place", "transforms": []}}
    field_defs = {
        "place": SimpleNamespace(field_type="vocab", label={"de": "Ort"}, is_required=False, is_repeatable=False),
    }
    result = dry_run(rows, mapping, field_defs)
    assert "place" in result["vocab_clusters"]
    cluster = result["vocab_clusters"]["place"][0]
    assert cluster["canonical"] == "Berlin"
    assert set(cluster["variants"]) == {"Berlin", "berlin", "Brlin"}


def test_dry_run_clusters_pending_vocab_field() -> None:
    """On-the-fly fields are merged into the dry-run as transient field defs with
    settings=None (never flushed). Clustering must tolerate that shape."""
    rows = [{"genre": v} for v in ("Electronic", "Electronic", "electronic", "Electronik")]
    mapping = {"genre": {"target": "genre", "transforms": []}}
    pending = SimpleNamespace(
        field_type="vocab", label={"de": "Genre"},
        is_required=False, is_repeatable=False, settings=None,
    )
    result = dry_run(rows, mapping, {"genre": pending})
    cluster = result["vocab_clusters"]["genre"][0]
    assert cluster["canonical"] == "Electronic"
    assert set(cluster["variants"]) == {"Electronic", "electronic", "Electronik"}
