import io
from typing import Iterator

from .base import Selector, SourceFormat, SourceRecord


class ExcelFormat(SourceFormat):
    """Excel (.xlsx) file format handler."""

    def sniff(self, content: bytes, filename: str) -> bool:
        """Match .xlsx/.xls files or Excel magic bytes."""
        lower = filename.lower()
        if lower.endswith((".xlsx", ".xls")):
            return True
        # Check for Excel magic bytes (PK signature for .xlsx)
        return content.startswith(b"PK\x03\x04") or content.startswith(b"\xd0\xcf\x11")

    def parse(self, content: bytes) -> Iterator[SourceRecord]:
        """Parse Excel workbook and yield rows as dicts."""
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            wb.close()
            return
        headers = [str(c) if c is not None else "" for c in header_row]
        for row in rows_iter:
            yield {
                headers[i]: str(v) if v is not None else ""
                for i, v in enumerate(row)
                if i < len(headers)
            }
        wb.close()

    def list_selectors(self, content: bytes) -> list[Selector]:
        """List headers as selectors."""
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        header_row = next(ws.iter_rows(values_only=True), None)
        wb.close()
        if not header_row:
            return []
        headers = [str(c) if c is not None else "" for c in header_row]
        return [
            Selector(path=h, label=h, sample="", kind="scalar")
            for h in headers
        ]
